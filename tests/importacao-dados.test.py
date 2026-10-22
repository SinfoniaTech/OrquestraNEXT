import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location("converter", ROOT/"scripts/import_orquestra.py")
converter=importlib.util.module_from_spec(spec);spec.loader.exec_module(converter)

class Importacao(unittest.TestCase):
 def test_reproducibilidade_sem_gravar(self):
  result=converter.convert(check=True)
  self.assertEqual(len(result),4)
 def test_valores_de_todos_os_campos_contra_excel(self):
  workbook=load_workbook(converter.SOURCE,read_only=True,data_only=True)
  try:
   for sheet in converter.SHEETS:
    raw=list(workbook[sheet].values);headers=raw[0]
    data=json.loads((converter.OUTPUT/(sheet+".json")).read_text())["records"]
    self.assertEqual(len(data),len(raw)-1)
    for original,record in zip(raw[1:],data):
     self.assertEqual(list(record),list(headers))
     for field,value in zip(headers,original):
      actual=record[field]
      if field in converter.DATES and value is not None:
       self.assertEqual(actual,value.strftime("%Y-%m-%dT%H:%M:%S")+"-03:00")
      else:
       self.assertEqual(actual,value,(sheet,record["ID_Chamado"],field))
  finally:workbook.close()
 def invalid(self,mutate,message):
  with tempfile.TemporaryDirectory(prefix="orquestra-invalid-") as directory:
   source=Path(directory)/"fixture.xlsx";output=Path(directory)/"output"
   wb=load_workbook(converter.SOURCE)
   mutate(wb);wb.save(source);wb.close()
   with self.assertRaisesRegex(ValueError,message):converter.convert(source,output)
   self.assertFalse(output.exists(),"Nenhum artefato parcial após erro de validação")
 def test_aba_ausente(self):
  self.invalid(lambda w:w.remove(w["T1_Acompanhamento"]),"Abas ausentes")
 def test_campo_ausente(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(1,46,"errado"),"cabeçalhos")
 def test_id_duplicado(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(3,1,w["HIST_60_DIAS"].cell(2,1).value),"duplicado")
 def test_tipo_numerico_invalido(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(2,33,"40"),"número")
 def test_categoria_invalida(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(2,41,"Azul"),"categoria")
 def test_data_invalida(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(2,28,"08/08/2026"),"data Excel")
 def test_formula_nao_materializada(self):
  self.invalid(lambda w:w["HIST_60_DIAS"].cell(2,33,"=40"),"fórmula")
if __name__=="__main__":unittest.main()
