import sys
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import import_feira_estendida as importer

class Feira(unittest.TestCase):
 def invalid(self,mutate,message):
  with tempfile.TemporaryDirectory() as temp:
   source=Path(temp)/'invalid.xlsx';output=Path(temp)/'generated.json'
   workbook=load_workbook(importer.SOURCE);mutate(workbook);workbook.save(source);workbook.close()
   with self.assertRaisesRegex(ValueError,message):importer.convert(source,output)
   self.assertFalse(output.exists())
 def test_determinismo_e_totais(self):
  result=importer.convert(check=True)
  self.assertEqual(len(result['snapshots']),12)
  self.assertEqual(len(result['events']),189)
 def test_id_duplicado(self):self.invalid(lambda w:w['FEIRA_ESTENDIDA'].cell(3,1,w['FEIRA_ESTENDIDA'].cell(2,1).value),'duplicados')
 def test_evento_futuro(self):self.invalid(lambda w:w['FEIRA_ESTENDIDA'].cell(2,29,w['FEIRA_ESTENDIDA'].cell(418,46).value),'Evento futuro')
 def test_auditoria_divergente(self):self.invalid(lambda w:w['FEIRA_EVENTOS'].cell(2,7,'Aberto'),'Auditoria')
 def test_criticidade_ativa(self):self.invalid(lambda w:w['FEIRA_ROTEIRO'].cell(2,11,99),'Indicador roteiro')
 def test_ordem(self):self.invalid(lambda w:w['FEIRA_ROTEIRO'].cell(2,1,'T01'),'Ordem')
 def test_numero_invalido(self):self.invalid(lambda w:w['FEIRA_ESTENDIDA'].cell(2,33,'21'),'número finito')
 def test_formula(self):self.invalid(lambda w:w['FEIRA_ESTENDIDA'].cell(2,33,'=21'),'Fórmula')
if __name__=='__main__':unittest.main()
