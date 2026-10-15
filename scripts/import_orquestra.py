"""Conversão controlada, sem alterações no XLSX. Requer openpyxl já disponível."""
import argparse
from collections import Counter
from datetime import datetime, date
import hashlib
import json
import math
from pathlib import Path
import unicodedata
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "datasets/ORQUESTRA_MVP_Base_60_Dias_Filtros_FINAL.xlsx"
OUTPUT = ROOT / "src/data/orquestra/generated"
ZONE = ZoneInfo("America/Sao_Paulo")
SHEETS = {"HIST_60_DIAS": (736, "2026-10-06T22:00:00-03:00"),
          "T0_Abertura": (10, "2026-09-08T08:00:00-03:00"),
          "T1_Acompanhamento": (16, "2026-09-08T08:30:00-03:00"),
          "T2_Desfecho": (24, "2026-09-08T09:30:00-03:00")}
DATES = {"Criado em", "Contactado", "Aceito", "Chegou ao Edifício", "Saída", "Momento_Dataset"}
INTEGERS = {"TA", "TB", "SLA_min", "ETA_min", "Recorrencia_30d"}
NUMBERS = INTEGERS | {"Percentual_SLA", "Custo_Estimado_R$"}
ENUMS = {"Status": {"Aberto", "Contactado", "Técnico a caminho", "Em atendimento", "Concluído"},
         "Criticidade_ORQUESTRA": {"Verde", "Amarelo", "Laranja", "Vermelho"},
         "Impacto_Contrato": {"Baixo", "Médio", "Alto"},
         **{k: {"Sim", "Não"} for k in ["VIP", "Corporativo", "Passageiro Preso", "Carro Parado"]}}

def normalize(key, value):
    if value is None or value == "":
        return None
    if key in DATES:
        if not isinstance(value, datetime) or value.tzinfo is not None:
            raise ValueError("esperada data Excel local sem fuso")
        return value.replace(tzinfo=ZONE).isoformat(timespec="seconds")
    if key in NUMBERS:
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError("esperado número finito")
        if key in INTEGERS and int(value) != value:
            raise ValueError("esperado inteiro")
        return int(value) if key in INTEGERS else value
    if not isinstance(value, str):
        raise ValueError("esperado texto")
    normalized = unicodedata.normalize("NFC", value.strip())
    return normalized or None

def convert(source=SOURCE, output=OUTPUT, check=False):
    schema = json.loads((ROOT / "src/data/orquestra/schema.json").read_text())
    headers = schema["fields"]
    assert len(headers) == 46 and len(set(headers)) == 46
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        missing = set(SHEETS) - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"Abas ausentes: {sorted(missing)}")
        digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
        artifacts = {}
        for sheet, (expected, reference) in SHEETS.items():
            ws = workbook[sheet]
            iterator = ws.iter_rows()
            head = next(iterator)
            if [c.value for c in head] != headers:
                raise ValueError(f"{sheet}: cabeçalhos diferentes dos 46 campos esperados")
            rows, ids, normalized_cells = [], set(), 0
            for line, cells in enumerate(iterator, 2):
                if all(c.value is None for c in cells):
                    continue
                if len(cells) != 46:
                    raise ValueError(f"{sheet}:{line}: quantidade de campos inválida")
                row = {}
                for key, cell in zip(headers, cells):
                    if cell.data_type == "f":
                        raise ValueError(f"{sheet}:{line}:{key}: fórmula não suportada; requer contrato")
                    try:
                        value = normalize(key, cell.value)
                        if key in ENUMS and value is not None and value not in ENUMS[key]:
                            raise ValueError("categoria desconhecida")
                        if key == "ID_Chamado" and value != cell.value:
                            raise ValueError("ID não pode ser alterado pela normalização")
                        if isinstance(cell.value, str) and value != cell.value:
                            normalized_cells += 1
                        row[key] = value
                    except ValueError as error:
                        raise ValueError(f"{sheet}:{line}:{key}: {error}") from error
                identity = row["ID_Chamado"]
                if not identity or identity in ids:
                    raise ValueError(f"{sheet}:{line}: ID ausente ou duplicado: {identity}")
                ids.add(identity)
                if row["Momento_Dataset"] != reference:
                    raise ValueError(f"{sheet}:{line}: momento de referência divergente")
                created = row["Criado em"]
                if created is None:
                    raise ValueError(f"{sheet}:{line}: Criado em obrigatório")
                if sheet == "HIST_60_DIAS" and not ("2026-08-08" <= created[:10] <= "2026-10-06" and created <= reference):
                    raise ValueError(f"{sheet}:{line}: criação fora da referência histórica")
                rows.append(row)
            if len(rows) != expected:
                raise ValueError(f"{sheet}: esperados {expected}, encontrados {len(rows)}")
            metadata = {"sourceFile": Path(source).name, "sourceSha256": digest,
                        "sheet": sheet, "kind": "history" if sheet == "HIST_60_DIAS" else "snapshot",
                        "timezone": "America/Sao_Paulo", "reference": reference, "synthetic": True,
                        "fieldCount": 46, "recordCount": len(rows), "schemaVersion": 1,
                        "statusCounts": dict(sorted(Counter(r["Status"] for r in rows).items())),
                        "criticalityCounts": dict(sorted(Counter(r["Criticidade_ORQUESTRA"] for r in rows).items())),
                        "nullCounts": {k: sum(r[k] is None for r in rows) for k in headers},
                        "normalizedTextCells": normalized_cells,
                        "normalization": "NFC/trim texto; números preservados; datas locais ISO com offset; sem recálculo"}
            artifacts[f"{sheet}.json"] = {"metadata": metadata, "records": rows}
        # Todas as fontes validadas antes de qualquer gravação.
        for filename, artifact in artifacts.items():
            text = json.dumps(artifact, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
            destination = Path(output) / filename
            if check:
                if not destination.exists() or destination.read_text() != text:
                    raise ValueError(f"Artefato desatualizado: {destination}")
            else:
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_text(text, encoding="utf-8")
        return {k: v["metadata"] for k, v in artifacts.items()}
    finally:
        workbook.close()

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true", help="validar sem gravar")
    args = parser.parse_args()
    print(json.dumps(convert(args.source, args.output, args.check), ensure_ascii=False, indent=2))
